import os
from flask import Flask, render_template, request, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename
from docx import Document
from deep_translator import GoogleTranslator
import time

app = Flask(__name__)
app.secret_key = 'super_secret_key'  # 用于 flash 消息

# 配置上传和下载文件夹
UPLOAD_FOLDER = 'uploads'
DOWNLOAD_FOLDER = 'downloads'
ALLOWED_EXTENSIONS = {'docx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['DOWNLOAD_FOLDER'] = DOWNLOAD_FOLDER

# 确保目录存在
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def translate_document(input_path, output_path):
    """
    读取 Word 文档，翻译内容，并保存到新文档
    """
    doc = Document(input_path)
    translator = GoogleTranslator(source='auto', target='zh-CN')
    
    # 创建新文档用于保存翻译结果
    new_doc = Document()
    
    total_paragraphs = len(doc.paragraphs)
    print(f"开始翻译，共 {total_paragraphs} 段...")

    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            try:
                # 限制文本长度以避免 API 错误（如果文本过长）
                text_to_translate = para.text[:4999] 
                translated_text = translator.translate(text_to_translate)
                new_doc.add_paragraph(translated_text)
                
                # 简单的进度打印
                if i % 10 == 0:
                    print(f"已处理 {i}/{total_paragraphs}")
                    
                # 避免请求过快
                # time.sleep(0.1) 
            except Exception as e:
                print(f"翻译段落 {i} 时出错: {e}")
                new_doc.add_paragraph(para.text) # 出错保留原文
        else:
            new_doc.add_paragraph("") # 保留空行

    new_doc.save(output_path)
    print("翻译完成！")

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        # 检查是否有文件部分
        if 'file' not in request.files:
            flash('没有文件被上传')
            return redirect(request.url)
        
        file = request.files['file']
        
        if file.filename == '':
            flash('未选择文件')
            return redirect(request.url)
            
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            input_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(input_path)
            
            # 生成输出文件名
            output_filename = f"translated_{filename}"
            output_path = os.path.join(app.config['DOWNLOAD_FOLDER'], output_filename)
            
            try:
                translate_document(input_path, output_path)
                return render_template('index.html', download_file=output_filename)
            except Exception as e:
                flash(f'翻译过程中发生错误: {str(e)}')
                return redirect(request.url)
        else:
            flash('仅支持 .docx 文件')
            return redirect(request.url)

    return render_template('index.html')

@app.route('/download/<filename>')
def download_file(filename):
    return send_file(os.path.join(app.config['DOWNLOAD_FOLDER'], filename), as_attachment=True)

if __name__ == '__main__':
    # 监听所有接口，方便外部访问（如果在云端环境）
    app.run(host='0.0.0.0', port=5000, debug=True)
